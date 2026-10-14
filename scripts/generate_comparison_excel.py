"""从当前 NL 与 Hybrid-P0-FIELDS-ONLY 汇总生成对比 Excel。"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def load_json(path: str | Path) -> dict:
    return json.loads(Path(path).read_text(encoding="utf-8"))


def _format_pct(value: float) -> str:
    return f"{value * 100:.1f}%"


def _fill_header(ws, row: int, cols: list[str]) -> None:
    fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    font = Font(bold=True)
    for col_idx, val in enumerate(cols, 1):
        cell = ws.cell(row=row, column=col_idx, value=val)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _write_row(ws, row: int, values: list, bold: bool = False) -> None:
    font = Font(bold=bold)
    for col_idx, val in enumerate(values, 1):
        cell = ws.cell(row=row, column=col_idx, value=val)
        cell.font = font
        cell.alignment = Alignment(
            horizontal="center" if col_idx > 1 else "left", vertical="center"
        )


def build_comparison(nl: dict, hybrid: dict) -> list[dict]:
    rows = [
        {"metric": "样本数", "nl": nl["n"], "hybrid": hybrid["n"]},
        {"metric": "正确数", "nl": nl["correct"], "hybrid": hybrid["correct"]},
        {
            "metric": "准确率",
            "nl": _format_pct(nl["accuracy"]),
            "hybrid": _format_pct(hybrid["accuracy"]),
        },
        {
            "metric": "平均通信 token",
            "nl": nl["avg_communication_tokens"],
            "hybrid": hybrid["avg_communication_tokens"],
        },
        {
            "metric": "平均输出 token",
            "nl": nl["avg_output_tokens"],
            "hybrid": hybrid["avg_output_tokens"],
        },
        {
            "metric": "通信 + 输出 token",
            "nl": round(nl["avg_communication_tokens"] + nl["avg_output_tokens"], 1),
            "hybrid": round(
                hybrid["avg_communication_tokens"] + hybrid["avg_output_tokens"], 1
            ),
        },
        {
            "metric": "平均模型调用次数",
            "nl": nl["avg_model_calls"],
            "hybrid": hybrid["avg_model_calls"],
        },
    ]
    return rows


def build_changes(rows: list[dict]) -> list[dict]:
    """为每一行补充 Hybrid-P0-FIELDS-ONLY 相对 NL 的变化。"""
    changed = []
    for r in rows:
        metric = r["metric"]
        nl_val = r["nl"]
        hybrid_val = r["hybrid"]

        def _delta_str(nl, p, is_pct=False):
            if isinstance(nl, (int, float)) and isinstance(p, (int, float)):
                delta = p - nl
                if is_pct:
                    return f"{delta:+.1f} pp"
                if nl == 0:
                    return "—"
                ratio = (p - nl) / nl
                return f"{delta:+.1f} ({ratio*100:+.1f}%)"
            return "—"

        is_acc = metric == "准确率"
        row = {
            "metric": metric,
            "nl": nl_val,
            "hybrid": hybrid_val,
            "hybrid_change": _delta_str(nl_val, hybrid_val, is_pct=is_acc),
        }
        changed.append(row)
    return changed


def write_excel(
    rows: list[dict],
    summaries: dict[str, dict],
    out_path: str | Path,
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "NL vs Hybrid 对比"

    headers = ["指标", "NL", "Hybrid-P0-FIELDS-ONLY", "Hybrid 相对 NL"]
    _fill_header(ws, 1, headers)

    for r_idx, row in enumerate(rows, 2):
        values = [
            row["metric"],
            row["nl"],
            row["hybrid"],
            row["hybrid_change"],
        ]
        _write_row(ws, r_idx, values)

    # 列宽自适应
    for col_idx in range(1, len(headers) + 1):
        max_len = len(headers[col_idx - 1])
        for row_idx in range(2, len(rows) + 2):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 4

    # 元数据 sheet
    meta_ws = wb.create_sheet("运行信息")
    meta_headers = ["条件", "run_dir", "protocol_path"]
    _fill_header(meta_ws, 1, meta_headers)
    for idx, (label, summary) in enumerate(summaries.items(), 2):
        _write_row(
            meta_ws,
            idx,
            [label, summary.get("run_dir", ""), summary.get("protocol_path", "")],
        )
    meta_ws.column_dimensions["A"].width = 16
    meta_ws.column_dimensions["B"].width = 60
    meta_ws.column_dimensions["C"].width = 60

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"[write] Excel saved to {out_path}")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Generate NL/Hybrid-P0-FIELDS-ONLY comparison Excel."
    )
    parser.add_argument(
        "--nl-summary",
        default="runs/hybrid_comm/nl_summary.json",
        help="Path to fresh NL summary JSON.",
    )
    parser.add_argument(
        "--hybrid-summary",
        default="runs/hybrid_comm/hybrid_preset_summary.json",
        help="Path to Hybrid-P0-FIELDS-ONLY summary JSON.",
    )
    parser.add_argument(
        "--out",
        default="runs/hybrid_comm/nl_hybrid_preset_comparison.xlsx",
        help="Output Excel path.",
    )
    args = parser.parse_args()

    nl_summary = load_json(args.nl_summary)
    hybrid_summary = load_json(args.hybrid_summary)

    rows = build_comparison(nl_summary, hybrid_summary)
    rows_with_changes = build_changes(rows)
    summaries = {"NL": nl_summary, "Hybrid-P0-FIELDS-ONLY": hybrid_summary}
    write_excel(rows_with_changes, summaries, args.out)


if __name__ == "__main__":
    main()
